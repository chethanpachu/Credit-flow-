import os
import tempfile
from datetime import datetime, date, timedelta
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_FILE = "grb_data.xlsx"

HEADERS = {
    "Areas": ["AreaID", "AreaName"],
    "Salesmen": ["SalesmanID", "SalesmanName", "PIN"],
    "Drivers": ["DriverID", "DriverName", "Phone"],
    "SalesmanAreas": ["SalesmanID", "AreaID"],
    "Bills": ["BillID", "BillNumber", "ShopName", "AreaID", "Amount", "Balance", "DateIssued", "Status"],
    "CollectionHistory": [
        "CollectionID", "BillID", "BillNumber", "ShopName", 
        "SalesmanID", "SalesmanName", "AmountCollected", 
        "CollectionDate", "RunningBalanceAfter", "VerificationStatus", "AdminNotes"
    ],
    "CashHandover": ["SalesmanID", "SalesmanName", "HandoverDate", "AmountHanded", "VerificationStatus", "AdminNotes"],
    "DriverAssignments": [
        "AssignmentID", "BillID", "BillNumber", "ShopName", "AreaID",
        "SalesmanID", "SalesmanName", "DriverID", "DriverName",
        "AmountToCollect", "AssignedDate", "ScheduledDate", "Status",
        "CollectedAmount", "Notes"
    ]
}

class ExcelManager:
    def __init__(self, file_path: str = EXCEL_FILE):
        self.file_path = file_path
        self.ensure_workbook_exists()

    def ensure_workbook_exists(self):
        """Creates workbook with all sheets and formatted headers if it does not exist, or adds missing sheets."""
        if not os.path.exists(self.file_path):
            wb = openpyxl.Workbook()
            default_sheet = wb.active
            wb.remove(default_sheet)

            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            align_center = Alignment(horizontal="center", vertical="center")

            for sheet_name, cols in HEADERS.items():
                ws = wb.create_sheet(title=sheet_name)
                ws.append(cols)
                for col_num in range(1, len(cols) + 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = align_center
                ws.row_dimensions[1].height = 24

            self._safe_save(wb)
            self._seed_default_data()
        else:
            # Upgrade existing workbook if any new sheet is missing
            wb = openpyxl.load_workbook(self.file_path)
            changed = False
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            align_center = Alignment(horizontal="center", vertical="center")

            for sheet_name, cols in HEADERS.items():
                if sheet_name not in wb.sheetnames:
                    ws = wb.create_sheet(title=sheet_name)
                    ws.append(cols)
                    for col_num in range(1, len(cols) + 1):
                        cell = ws.cell(row=1, column=col_num)
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = align_center
                    ws.row_dimensions[1].height = 24
                    changed = True

            if changed:
                self._safe_save(wb)

    def _safe_save(self, wb: openpyxl.Workbook):
        """Atomic safe save: writes to a temporary file then replaces target to prevent corruption."""
        dir_name = os.path.dirname(os.path.abspath(self.file_path)) or "."
        temp_file = tempfile.NamedTemporaryFile(delete=False, dir=dir_name, suffix=".tmp")
        temp_path = temp_file.name
        temp_file.close()

        try:
            wb.save(temp_path)
            if os.path.exists(self.file_path):
                os.replace(temp_path, self.file_path)
            else:
                os.rename(temp_path, self.file_path)
        except Exception as e:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            raise e

    def _seed_default_data(self):
        """Seeds standard 4 salesmen, drivers, areas, and sample bills."""
        areas = [
            ("A1", "Market Yard"),
            ("A2", "Gandhi Bazar"),
            ("A3", "Malleshwaram"),
            ("A4", "Rajajinagar"),
            ("A5", "Jayanagar"),
            ("A6", "Basavanagudi"),
            ("A7", "Shivajinagar"),
            ("A8", "Indiranagar"),
        ]
        salesmen = [
            ("S1", "Ramesh Kumar", "1234"),
            ("S2", "Suresh Patel", "1234"),
            ("S3", "Ganesh Rao", "1234"),
            ("S4", "Mahesh Gowda", "1234"),
        ]
        drivers = [
            ("D1", "Anand Swamy (Van #KA-04-1029)", "9845012345"),
            ("D2", "Manjunath K (Van #KA-05-8842)", "9845067890"),
        ]
        salesman_areas = [
            ("S1", "A1"), ("S1", "A2"),
            ("S2", "A3"), ("S2", "A4"),
            ("S3", "A5"), ("S3", "A6"),
            ("S4", "A7"), ("S4", "A8"),
        ]

        wb = openpyxl.load_workbook(self.file_path)
        for a in areas:
            wb["Areas"].append(list(a))
        for s in salesmen:
            wb["Salesmen"].append(list(s))
        for d in drivers:
            wb["Drivers"].append(list(d))
        for sa in salesman_areas:
            wb["SalesmanAreas"].append(list(sa))

        today = date.today()
        d_recent = (today - timedelta(days=5)).strftime("%Y-%m-%d")
        d_overdue = (today - timedelta(days=38)).strftime("%Y-%m-%d")
        tomorrow = (today + timedelta(days=1)).strftime("%Y-%m-%d")
        
        sample_bills = [
            ("B-1001", "GRB-801", "Annapoorna Provision", "A1", 12500, 12500, d_recent, "pending"),
            ("B-1002", "GRB-802", "Sri Krishna Sweets & Ghee", "A2", 8400, 8400, d_recent, "pending"),
            ("B-1003", "GRB-803", "Mahalakshmi Stores", "A3", 16200, 16200, d_recent, "pending"),
            ("B-1004", "GRB-804", "Venkateshwara Traders", "A4", 9500, 9500, d_recent, "pending"),
            ("B-1005", "GRB-805", "Balaji Kondapalli Stores", "A5", 14000, 14000, d_overdue, "pending"),
            ("B-1006", "GRB-806", "Cauvery Supermarket", "A6", 21000, 0, (today - timedelta(days=12)).strftime("%Y-%m-%d"), "paid"),
        ]
        for b in sample_bills:
            wb["Bills"].append(list(b))

        wb["CollectionHistory"].append([
            "COL-501", "B-1006", "GRB-806", "Cauvery Supermarket", 
            "S3", "Ganesh Rao", 21000, (today - timedelta(days=2)).strftime("%Y-%m-%d"), 0, "Verified", "Amount matched voucher"
        ])

        # Sample initial driver assignment for tomorrow
        wb["DriverAssignments"].append([
            "DRV-101", "B-1001", "GRB-801", "Annapoorna Provision", "A1",
            "S1", "Ramesh Kumar", "D1", "Anand Swamy (Van #KA-04-1029)",
            12500, today.strftime("%Y-%m-%d"), tomorrow, "Assigned", 0, "Collect payment before 11 AM"
        ])

        self._safe_save(wb)

    def _read_sheet_dicts(self, sheet_name: str) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(self.file_path, data_only=True)
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            return []
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        result = []
        for row in rows[1:]:
            if all(v is None for v in row):
                continue
            item = {}
            for col_idx, h in enumerate(headers):
                item[h] = row[col_idx] if col_idx < len(row) else None
            result.append(item)
        return result

    # --- Setup Methods ---
    def get_areas(self) -> List[Dict[str, Any]]:
        return self._read_sheet_dicts("Areas")

    def add_area(self, area_id: str, area_name: str) -> None:
        area_id = area_id.strip().upper()
        area_name = area_name.strip()
        if not area_id or not area_name:
            raise ValueError("Area ID and Area Name cannot be empty.")
        areas = self.get_areas()
        if any(a.get("AreaID") == area_id for a in areas):
            raise ValueError(f"Area ID '{area_id}' already exists.")
        wb = openpyxl.load_workbook(self.file_path)
        wb["Areas"].append([area_id, area_name])
        self._safe_save(wb)

    def get_salesmen(self) -> List[Dict[str, Any]]:
        return self._read_sheet_dicts("Salesmen")

    def add_salesman(self, salesman_id: str, salesman_name: str, pin: str = "1234") -> None:
        salesman_id = salesman_id.strip().upper()
        salesman_name = salesman_name.strip()
        pin = pin.strip() or "1234"
        if not salesman_id or not salesman_name:
            raise ValueError("Salesman ID and Salesman Name cannot be empty.")
        salesmen = self.get_salesmen()
        if any(s.get("SalesmanID") == salesman_id for s in salesmen):
            raise ValueError(f"Salesman ID '{salesman_id}' already exists.")
        wb = openpyxl.load_workbook(self.file_path)
        wb["Salesmen"].append([salesman_id, salesman_name, pin])
        self._safe_save(wb)

    def get_drivers(self) -> List[Dict[str, Any]]:
        return self._read_sheet_dicts("Drivers")

    def add_driver(self, driver_id: str, driver_name: str, phone: str = "") -> None:
        driver_id = driver_id.strip().upper()
        driver_name = driver_name.strip()
        if not driver_id or not driver_name:
            raise ValueError("Driver ID and Name cannot be empty.")
        drivers = self.get_drivers()
        if any(d.get("DriverID") == driver_id for d in drivers):
            raise ValueError(f"Driver ID '{driver_id}' already exists.")
        wb = openpyxl.load_workbook(self.file_path)
        wb["Drivers"].append([driver_id, driver_name, phone.strip()])
        self._safe_save(wb)

    def get_salesman_areas(self) -> List[Dict[str, Any]]:
        return self._read_sheet_dicts("SalesmanAreas")

    def get_areas_for_salesman(self, salesman_id: str) -> List[Dict[str, Any]]:
        sa_list = self.get_salesman_areas()
        assigned_ids = {sa.get("AreaID") for sa in sa_list if str(sa.get("SalesmanID")) == str(salesman_id)}
        all_areas = self.get_areas()
        return [a for a in all_areas if a.get("AreaID") in assigned_ids]

    def set_salesman_areas(self, salesman_id: str, area_ids: List[str]) -> None:
        salesman_id = salesman_id.strip()
        wb = openpyxl.load_workbook(self.file_path)
        ws = wb["SalesmanAreas"]
        rows = list(ws.iter_rows(values_only=True))
        headers = rows[0] if rows else HEADERS["SalesmanAreas"]
        new_rows = [headers]
        if len(rows) > 1:
            for r in rows[1:]:
                if r and r[0] != salesman_id:
                    new_rows.append(r)
        for aid in area_ids:
            new_rows.append((salesman_id, aid))
        wb.remove(ws)
        new_ws = wb.create_sheet("SalesmanAreas")
        for r in new_rows:
            new_ws.append(list(r))
        self._safe_save(wb)

    # --- Bills Methods ---
    def get_bills(self) -> List[Dict[str, Any]]:
        return self._read_sheet_dicts("Bills")

    def check_double_billing(self, shop_name: str) -> Optional[Dict[str, Any]]:
        """A customer with an unpaid old bill cannot receive a new bill until the old one is cleared."""
        shop_norm = shop_name.strip().lower()
        for b in self.get_bills():
            if str(b.get("ShopName") or "").strip().lower() == shop_norm:
                status = str(b.get("Status") or "").strip().lower()
                balance = float(b.get("Balance") or 0)
                if status == "pending" or balance > 0:
                    return b
        return None

    def add_bill(self, area_id: str, bill_number: str, shop_name: str, amount: float, date_issued: str) -> Dict[str, Any]:
        """Area selection comes first."""
        area_id = area_id.strip()
        bill_number = bill_number.strip()
        shop_name = shop_name.strip()
        if not area_id:
            raise ValueError("Area must be selected first.")
        if not bill_number or not shop_name or amount <= 0:
            raise ValueError("All fields are required and amount must be greater than zero.")

        # 1. Block duplicate Bill Numbers
        for b in self.get_bills():
            if str(b.get("BillNumber") or "").strip().lower() == bill_number.lower():
                raise ValueError(f"Duplicate Bill Number! Bill '{bill_number}' already exists in system.")

        # 2. Block Double Billing
        pending_bill = self.check_double_billing(shop_name)
        if pending_bill:
            pending_num = pending_bill.get("BillNumber")
            pending_bal = pending_bill.get("Balance")
            pending_dt = pending_bill.get("DateIssued")
            raise ValueError(
                f"Double Billing Blocked! Shop '{shop_name}' already has an unpaid pending bill "
                f"#{pending_num} of ₹{pending_bal:,.2f} issued on {pending_dt}. "
                f"The old bill must be fully cleared before issuing a new bill."
            )

        bills = self.get_bills()
        bill_id = f"B-{len(bills) + 1001}"
        new_row = [bill_id, bill_number, shop_name, area_id, amount, amount, date_issued, "pending"]
        wb = openpyxl.load_workbook(self.file_path)
        wb["Bills"].append(new_row)
        self._safe_save(wb)

        return {
            "BillID": bill_id, "BillNumber": bill_number, "ShopName": shop_name,
            "AreaID": area_id, "Amount": amount, "Balance": amount,
            "DateIssued": date_issued, "Status": "pending"
        }

    # --- Evening Collection Methods ---
    def get_pending_bills_for_salesman(self, salesman_id: str) -> List[Dict[str, Any]]:
        """Privacy: only bills in this salesman's assigned areas."""
        sa_list = self.get_salesman_areas()
        assigned_areas = {sa.get("AreaID") for sa in sa_list if str(sa.get("SalesmanID")) == str(salesman_id)}
        pending = []
        for b in self.get_bills():
            status = str(b.get("Status") or "").strip().lower()
            balance = float(b.get("Balance") or 0)
            area_id = str(b.get("AreaID") or "").strip()
            if (status == "pending" or balance > 0) and (area_id in assigned_areas):
                pending.append(b)
        return pending

    def record_collection(self, bill_number: str, salesman_id: str, amount_collected: float, collection_date: str) -> Dict[str, Any]:
        if amount_collected <= 0:
            raise ValueError("Collection amount must be > 0.")
        salesmen = self.get_salesmen()
        s_obj = next((s for s in salesmen if str(s.get("SalesmanID")) == str(salesman_id)), None)
        salesman_name = s_obj.get("SalesmanName") if s_obj else "Unknown"

        wb = openpyxl.load_workbook(self.file_path)
        ws_bills = wb["Bills"]
        bill_row_idx = None
        bill_data = None
        for idx, row in enumerate(ws_bills.iter_rows(values_only=True), start=1):
            if idx == 1:
                continue
            if row and str(row[1]).strip() == bill_number.strip():
                bill_row_idx = idx
                bill_data = {
                    "BillID": row[0], "BillNumber": row[1], "ShopName": row[2],
                    "Balance": float(row[5] or 0)
                }
                break

        if not bill_row_idx or not bill_data:
            raise ValueError(f"Bill '{bill_number}' not found.")

        current_balance = bill_data["Balance"]
        if amount_collected > current_balance:
            raise ValueError(f"Amount ₹{amount_collected:,.2f} exceeds current balance of ₹{current_balance:,.2f}!")

        new_balance = round(current_balance - amount_collected, 2)
        new_status = "paid" if new_balance <= 0 else "pending"

        ws_bills.cell(row=bill_row_idx, column=6, value=new_balance)
        ws_bills.cell(row=bill_row_idx, column=8, value=new_status)

        ws_coll = wb["CollectionHistory"]
        coll_id = f"COL-{ws_coll.max_row + 100}"
        ws_coll.append([
            coll_id, bill_data["BillID"], bill_data["BillNumber"], bill_data["ShopName"],
            salesman_id, salesman_name, amount_collected, collection_date, new_balance,
            "Pending Review", ""
        ])
        self._safe_save(wb)

        return {
            "CollectionID": coll_id, "BillNumber": bill_number,
            "AmountCollected": amount_collected, "RemainingBalance": new_balance,
            "NewStatus": new_status, "CollectionDate": collection_date
        }

    # --- Driver Assignment Methods ---
    def assign_bill_to_driver(
        self, bill_number: str, salesman_id: str, driver_id: str, 
        amount_to_collect: float, scheduled_date: str, notes: str = ""
    ) -> Dict[str, Any]:
        """Assigns today's bill collection to driver for next day."""
        bills = self.get_bills()
        bill = next((b for b in bills if str(b.get("BillNumber")) == str(bill_number)), None)
        if not bill:
            raise ValueError(f"Bill #{bill_number} not found.")

        salesmen = self.get_salesmen()
        s_obj = next((s for s in salesmen if str(s.get("SalesmanID")) == str(salesman_id)), None)
        salesman_name = s_obj.get("SalesmanName") if s_obj else "Unknown"

        drivers = self.get_drivers()
        d_obj = next((d for d in drivers if str(d.get("DriverID")) == str(driver_id)), None)
        driver_name = d_obj.get("DriverName") if d_obj else "Unknown"

        wb = openpyxl.load_workbook(self.file_path)
        ws = wb["DriverAssignments"]
        assignment_id = f"DRV-{ws.max_row + 100}"
        today_str = date.today().strftime("%Y-%m-%d")

        row = [
            assignment_id,
            bill.get("BillID"),
            bill.get("BillNumber"),
            bill.get("ShopName"),
            bill.get("AreaID"),
            salesman_id,
            salesman_name,
            driver_id,
            driver_name,
            amount_to_collect,
            today_str,
            scheduled_date,
            "Assigned",
            0,
            notes.strip()
        ]
        ws.append(row)
        self._safe_save(wb)

        return {
            "AssignmentID": assignment_id,
            "BillNumber": bill_number,
            "ShopName": bill.get("ShopName"),
            "DriverName": driver_name,
            "ScheduledDate": scheduled_date,
            "AmountToCollect": amount_to_collect
        }

    def get_driver_assignments(self, salesman_id: Optional[str] = None) -> List[Dict[str, Any]]:
        all_assign = self._read_sheet_dicts("DriverAssignments")
        if salesman_id:
            # Salesman privacy: only assignments created by this salesman
            return [a for a in all_assign if str(a.get("SalesmanID")) == str(salesman_id)]
        return all_assign

    def update_driver_assignment_status(self, assignment_id: str, status: str, collected_amount: float = 0.0) -> None:
        wb = openpyxl.load_workbook(self.file_path)
        ws = wb["DriverAssignments"]
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx > 1 and row and str(row[0]) == str(assignment_id):
                ws.cell(row=idx, column=13, value=status) # Status
                ws.cell(row=idx, column=14, value=collected_amount) # CollectedAmount
                break
        self._safe_save(wb)

    # --- Admin Verification & Audit Methods ---
    def verify_collection(self, collection_id: str, status: str, admin_notes: str = "") -> None:
        """Admin marks an evening collection entry as Verified / Rejected."""
        wb = openpyxl.load_workbook(self.file_path)
        ws = wb["CollectionHistory"]
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx > 1 and row and str(row[0]) == str(collection_id):
                ws.cell(row=idx, column=10, value=status) # VerificationStatus
                ws.cell(row=idx, column=11, value=admin_notes) # AdminNotes
                break
        self._safe_save(wb)

    def verify_cash_handover(self, salesman_id: str, handover_date: str, status: str, admin_notes: str = "") -> None:
        """Admin marks cash handover as Verified / Discrepancy."""
        wb = openpyxl.load_workbook(self.file_path)
        ws = wb["CashHandover"]
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx > 1 and row and str(row[0]) == str(salesman_id) and str(row[2]) == str(handover_date):
                ws.cell(row=idx, column=5, value=status)
                ws.cell(row=idx, column=6, value=admin_notes)
                break
        self._safe_save(wb)

    def get_collections_by_date(self, target_date: str, salesman_id: Optional[str] = None) -> List[Dict[str, Any]]:
        all_coll = self._read_sheet_dicts("CollectionHistory")
        filtered = [c for c in all_coll if str(c.get("CollectionDate") or "").strip() == target_date.strip()]
        if salesman_id:
            return [c for c in filtered if str(c.get("SalesmanID")) == str(salesman_id)]
        return filtered

    def get_cash_handovers(self, salesman_id: Optional[str] = None) -> List[Dict[str, Any]]:
        all_h = self._read_sheet_dicts("CashHandover")
        if salesman_id:
            return [h for h in all_h if str(h.get("SalesmanID")) == str(salesman_id)]
        return all_h

    def save_cash_handover(self, salesman_id: str, handover_date: str, amount_handed: float) -> None:
        salesmen = self.get_salesmen()
        s_obj = next((s for s in salesmen if str(s.get("SalesmanID")) == str(salesman_id)), None)
        salesman_name = s_obj.get("SalesmanName") if s_obj else "Unknown"

        wb = openpyxl.load_workbook(self.file_path)
        ws = wb["CashHandover"]
        existing_row = None
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx > 1 and row and str(row[0]) == str(salesman_id) and str(row[2]) == str(handover_date):
                existing_row = idx
                break

        if existing_row:
            ws.cell(row=existing_row, column=4, value=amount_handed)
            ws.cell(row=existing_row, column=5, value="Pending Review")
        else:
            ws.append([salesman_id, salesman_name, handover_date, amount_handed, "Pending Review", ""])
        self._safe_save(wb)

    def get_reconciliation_summary(self, target_date: str) -> List[Dict[str, Any]]:
        salesmen = self.get_salesmen()
        collections_today = self.get_collections_by_date(target_date)
        handovers = self.get_cash_handovers()
        summary = []
        for s in salesmen:
            s_id = str(s.get("SalesmanID"))
            s_name = str(s.get("SalesmanName"))
            s_colls = [c for c in collections_today if str(c.get("SalesmanID")) == s_id]
            system_total = sum(float(c.get("AmountCollected") or 0) for c in s_colls)
            handover_obj = next((h for h in handovers if str(h.get("SalesmanID")) == s_id and str(h.get("HandoverDate")) == target_date), None)
            amount_handed = float(handover_obj.get("AmountHanded") or 0) if handover_obj else 0.0
            v_status = handover_obj.get("VerificationStatus") if handover_obj else "Not Handed"
            diff = round(amount_handed - system_total, 2)
            summary.append({
                "SalesmanID": s_id, "SalesmanName": s_name, "SystemTotal": system_total,
                "AmountHanded": amount_handed, "Difference": diff, "IsMatched": abs(diff) < 0.01,
                "VerificationStatus": v_status,
                "Collections": s_colls
            })
        return summary
